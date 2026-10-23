"""Write the timetable workbook: long-format Schedule, one grid per section, one grid per faculty, Validation."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from common import DAYS, LUNCH

HEAD = PatternFill("solid", fgColor="1F3864")
LABF = PatternFill("solid", fgColor="E2EFDA")
LUNF = PatternFill("solid", fgColor="D9D9D9")
thin = Side(style="thin", color="999999")


def _grid(ws, df, title, cell_text):
    ws["A1"] = title
    ws["A1"].font = Font(name="Arial", bold=True, size=12)
    ws.append([])
    ws.append([""] + [f"S{p}" for p in range(1, 10)])
    for c in ws[3]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = HEAD
        c.alignment = Alignment(horizontal="center")
    for d in DAYS:
        row = [d]
        for p in range(1, 10):
            if p == LUNCH:
                row.append("LUNCH"); continue
            m = df[(df["day"] == d) & (df["period"] == p)]
            row.append("\n".join(cell_text(r) for r in m.itertuples()))
        ws.append(row)
    for r in ws.iter_rows(min_row=4, max_row=8, max_col=10):
        ws.row_dimensions[r[0].row].height = 36
        for c in r:
            c.font = Font(name="Arial", size=9, bold=(c.column == 1))
            c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if c.value == "LUNCH": c.fill = LUNF
            elif c.column > 1 and c.value and "LAB" in str(c.value): c.fill = LABF
    ws.column_dimensions["A"].width = 7
    for col in "BCDEFGHIJ":
        ws.column_dimensions[col].width = 17


def write_workbook(path, sched, violations=None, summary=None):
    wb = Workbook()
    ws = wb.active; ws.title = "Schedule"
    ws.append(list(sched.columns))
    for r in sched.itertuples(index=False):
        ws.append(list(r))
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = HEAD
    for col, w in zip("ABCDEFGH", (9, 7, 8, 9, 11, 9, 8, 22)):
        ws.column_dimensions[col].width = w
    for sec in sorted(sched["section"].unique()):
        g = sched[sched["section"] == sec]
        _grid(wb.create_sheet(f"Section_{sec}"), g, f"Section {sec} (synthetic)",
              lambda r: f"{r.subject}{' LAB' if r.type == 'Lab' else ''}\n{r.faculty_id} {r.room}")
    for f in sorted(sched["faculty_id"].unique()):
        g = sched[sched["faculty_id"] == f]
        _grid(wb.create_sheet(f"Faculty_{f}"), g, f"Faculty {f}: {len(g)} h/week",
              lambda r: f"{r.section}-{r.subject}{' LAB' if r.type == 'Lab' else ''}\n{r.room}")
    wv = wb.create_sheet("Validation")
    if summary:
        wv.append(["Summary", summary])
    wv.append(["check", "severity", "detail"])
    if violations is None or violations.empty:
        wv.append(["ALL_CHECKS", "PASS", "No clashes or rule violations found"])
    else:
        for r in violations.itertuples(index=False):
            wv.append(list(r))
    wv.column_dimensions["A"].width = 26; wv.column_dimensions["C"].width = 90
    wb.save(path)
